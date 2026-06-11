import os
import subprocess
import asyncio
from pathlib import Path

import docx
import openpyxl
import xlrd

from app.services.ocr_service import extract_text_with_tesseract


def _extract_docx(file_path: str) -> str:
    doc = docx.Document(file_path)
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)


def _extract_doc(file_path: str) -> str:
    # Run antiword to extract text from legacy .doc files
    try:
        result = subprocess.run(["antiword", file_path], capture_output=True, text=True, check=True)
        return result.stdout
    except subprocess.CalledProcessError as exc:
        raise RuntimeError(f"Failed to extract .doc using antiword: {exc.stderr}") from exc


def _extract_xlsx(file_path: str) -> str:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_data = [str(cell) for cell in row if cell is not None]
            if row_data:
                lines.append("\t".join(row_data))
    return "\n".join(lines)


def _extract_xls(file_path: str) -> str:
    wb = xlrd.open_workbook(file_path)
    lines = []
    for sheet in wb.sheets():
        for row_idx in range(sheet.nrows):
            row_data = [str(cell.value) for cell in sheet.row(row_idx) if cell.value not in ("", None)]
            if row_data:
                lines.append("\t".join(row_data))
    return "\n".join(lines)


async def extract_text_from_file(file_path: str) -> str:
    """Extract raw text from a document based on its extension."""
    ext = Path(file_path).suffix.lower()
    
    if ext == ".docx":
        return await asyncio.to_thread(_extract_docx, file_path)
    elif ext == ".doc":
        return await asyncio.to_thread(_extract_doc, file_path)
    elif ext == ".xlsx":
        return await asyncio.to_thread(_extract_xlsx, file_path)
    elif ext == ".xls":
        return await asyncio.to_thread(_extract_xls, file_path)
    
    # Fallback to Tesseract OCR for PDFs and images
    return await extract_text_with_tesseract(file_path)
